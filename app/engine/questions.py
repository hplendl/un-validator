"""Stakeholder questions generated from open findings.

One question per check, layer and severity, so a class with many similar defects
becomes a single question linked to every finding id. Nothing here changes a score.
"""

from __future__ import annotations

import csv
from pathlib import Path

from .models import STAGE_NAMES

ROLES = {
    "metadata": "Data steward",
    "lineage": "GIS administrator",
    "schema": "Data modeler",
    "quality": "Data editor",
    "discover": "GIS administrator",
    "report": "Data steward",
}
PRIORITY = {"error": "high", "warning": "medium", "info": "low"}
WHY = {
    "metadata": "Incomplete metadata means the next person cannot tell what the data is, who owns it, or how it may be used.",
    "lineage": "Without a clear source and coordinate system, a later load can shift or mis-credit the data.",
    "schema": "A class that does not match the reference model will fail or silently drop attributes when it is loaded.",
    "quality": "Records that fail this check are unsafe to treat as trusted until someone confirms or corrects them.",
    "discover": "If the dataset did not open completely, every later conclusion about it is incomplete.",
    "report": "This note is part of the published result and should be confirmed or closed.",
}


def _question_text(finding: dict) -> str:
    layer = finding.get("layer") or "the dataset"
    check = finding.get("check") or "this check"
    return f"Is the '{check}' result on {layer} expected, and if not, who will correct it?"


def build_questions(result: dict, limit: int = 200) -> list[dict]:
    """Open questions for findings that are still OPEN. Grouped so the list stays readable."""
    groups: dict[tuple, dict] = {}
    for ds in result.get("datasets") or []:
        for f in ds.get("findings") or []:
            if str(f.get("disposition") or "OPEN").upper() != "OPEN":
                continue
            key = (ds.get("dataset"), f.get("stage"), f.get("check"), f.get("layer") or "", f.get("severity"))
            g = groups.get(key)
            if g is None:
                stage = f.get("stage") or ""
                g = {
                    "dataset": ds.get("dataset") or "",
                    "question": _question_text(f),
                    "why_it_matters": WHY.get(stage, WHY["report"]),
                    "role": ROLES.get(stage, "Data steward"),
                    "priority": PRIORITY.get(f.get("severity") or "", "low"),
                    "stage": STAGE_NAMES.get(stage, stage),
                    "check": f.get("check") or "",
                    "layer": f.get("layer") or "",
                    "severity": f.get("severity") or "",
                    "finding_ids": [],
                    "status": "OPEN",
                }
                groups[key] = g
            fid = f.get("finding_id")
            if fid and fid not in g["finding_ids"]:
                g["finding_ids"].append(fid)
    order = {"high": 0, "medium": 1, "low": 2}
    rows = sorted(groups.values(), key=lambda r: (order.get(r["priority"], 9), r["dataset"], r["check"], r["layer"]))
    return rows[:limit]


def write_questions_csv(rows: list[dict], path: Path) -> None:
    from .report import csv_safe

    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(
            [
                "dataset",
                "priority",
                "role",
                "status",
                "stage",
                "check",
                "layer",
                "question",
                "why_it_matters",
                "finding_ids",
            ]
        )
        for r in rows:
            w.writerow(
                [
                    csv_safe(r["dataset"]),
                    r["priority"],
                    csv_safe(r["role"]),
                    r["status"],
                    csv_safe(r["stage"]),
                    csv_safe(r["check"]),
                    csv_safe(r["layer"]),
                    csv_safe(r["question"]),
                    csv_safe(r["why_it_matters"]),
                    " ".join(r["finding_ids"]),
                ]
            )


def write_questions_xlsx(rows: list[dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    headers = [
        "dataset",
        "priority",
        "role",
        "status",
        "stage",
        "check",
        "layer",
        "question",
        "why_it_matters",
        "finding_ids",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(
            [
                r["dataset"],
                r["priority"],
                r["role"],
                r["status"],
                r["stage"],
                r["check"],
                r["layer"],
                r["question"],
                r["why_it_matters"],
                " ".join(r["finding_ids"]),
            ]
        )
    if rows:
        ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        ws.add_table(
            Table(
                displayName="Questions",
                ref=ref,
                tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=True),
            )
        )
        dv = DataValidation(type="list", formula1='"OPEN,CLOSED,ACCEPTED"', allow_blank=False)
        dv.add(f"D2:D{len(rows) + 1}")
        ws.add_data_validation(dv)
    widths = [28, 12, 22, 12, 28, 32, 24, 64, 64, 44]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    wb.save(path)
